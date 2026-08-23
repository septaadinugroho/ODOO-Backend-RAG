"""Corpus ingest: read the CSV/XLSX, validate V1–V8, load into PostgreSQL.

Column list and validation rules come from the Data Dictionary. The DD is
emphatic that validation happens *before* anything reaches a table, and that the
rules have different blast radii, so they are modelled explicitly rather than
collapsed into one pass/fail:

    reject_batch   V1        — nothing is written; the file is wrong
    reject_row     V2–V5     — that record is dropped, the rest load
    warning        V6, V7    — loaded, flagged for manual review
    hard_warning   V8        — loaded, but a v1 design assumption just moved

V8 is inverted on purpose. Everywhere else an empty field is the problem; here a
*populated* fee_rm or schedule is the alarm, because the whole v1 API contract is
built on those two columns being empty (FR-16, FR-17). If they ever arrive with
values, someone has to make a conscious decision about the contract — so the
loader shouts instead of quietly accepting the data.
"""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import TYPE_CHECKING, Dict, Iterable, List, Optional, Sequence, Tuple

from app.core.vocabulary import (
    CURRICULA,
    DELIVERY_MODES,
    EDUCATION_STAGES,
    INSTRUCTION_LANGUAGES,
    LEVEL_SORT_ORDER,
    LEVEL_STAGE,
    LEVELS,
    MODE_TO_GROUP,
    SCHOOL_CATEGORIES,
    STATES,
    SUBJECTS,
)

if TYPE_CHECKING:  # pragma: no cover
    from sqlalchemy.orm import Session

# The ORM is imported lazily inside the loading functions, so that reading and
# validating a corpus file needs nothing but the standard library. That keeps
# `python scripts/load_data.py --dry-run` usable before a database exists, and
# lets the V1–V8 tests run without a live engine.

# ─── The 23 columns, per the Data Dictionary ────────────────────────────────

REQUIRED_COLUMNS: Tuple[str, ...] = (
    "record_id",
    "class_title",
    "subject",
    "level",
    "state",
    "delivery_mode",
    "instruction_language",
    "curriculum",
    "education_stage",
    "school_category",
    "city_area",
    "provider_name",
    "provider_source_url",
    "fee_rm",
    "schedule",
    "verified_on",
    "class_status",
    "data_label",
    "source_basis",
    "curriculum_source_url",
    "retrieval_text",
    "keywords",
    "random_seed",
)

# DD: "random_seed — Artefak generasi (20260818). Tidak dimuat ke produksi."
NOT_LOADED: Tuple[str, ...] = ("random_seed",)

# V5: the eleven FR-12 fields, as they appear in the CSV.
V5_REQUIRED_NON_EMPTY: Tuple[str, ...] = (
    "record_id",
    "class_title",
    "provider_name",
    "level",
    "subject",
    "delivery_mode",
    "instruction_language",
    "state",
    "city_area",
    "provider_source_url",
    "verified_on",
)

# V2: which CSV column is checked against which closed list.
ENUM_COLUMNS: Dict[str, Sequence[str]] = {
    "subject": SUBJECTS,
    "level": LEVELS,
    "state": STATES,
    "delivery_mode": DELIVERY_MODES,
    "instruction_language": INSTRUCTION_LANGUAGES,
    "curriculum": CURRICULA,
    "education_stage": EDUCATION_STAGES,
    "school_category": SCHOOL_CATEGORIES,
}

_RECORD_ID = re.compile(r"^MYTC-\d{3}$")
_HTTPS_URI = re.compile(r"^https://[^\s]+$")
_ISO_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")

REJECT_BATCH = "reject_batch"
REJECT_ROW = "reject_row"
WARNING = "warning"
HARD_WARNING = "hard_warning"


# ─── Reporting ──────────────────────────────────────────────────────────────


@dataclass
class Issue:
    rule: str
    severity: str
    message: str
    record_id: Optional[str] = None
    column: Optional[str] = None

    def __str__(self) -> str:
        where = f" [{self.record_id}" + (f".{self.column}]" if self.column else "]")
        return f"{self.rule} {self.severity}{where if self.record_id else ''}: {self.message}"


@dataclass
class IngestReport:
    source: str = ""
    rows_read: int = 0
    rows_loaded: int = 0
    rows_rejected: int = 0
    batch_rejected: bool = False
    issues: List[Issue] = field(default_factory=list)
    unknown_enum_values: Dict[str, List[str]] = field(default_factory=dict)

    def add(
        self,
        rule: str,
        severity: str,
        message: str,
        record_id: Optional[str] = None,
        column: Optional[str] = None,
    ) -> None:
        self.issues.append(Issue(rule, severity, message, record_id, column))
        if severity == REJECT_BATCH:
            self.batch_rejected = True

    def by_severity(self, severity: str) -> List[Issue]:
        return [i for i in self.issues if i.severity == severity]

    @property
    def ok(self) -> bool:
        return not self.batch_rejected and self.rows_loaded > 0

    def summary(self) -> str:
        lines = [
            f"source          : {self.source}",
            f"rows read       : {self.rows_read}",
            f"rows loaded     : {self.rows_loaded}",
            f"rows rejected   : {self.rows_rejected}",
            f"batch rejected  : {self.batch_rejected}",
        ]
        for label, severity in (
            ("BATCH REJECT", REJECT_BATCH),
            ("ROW REJECT", REJECT_ROW),
            ("HARD WARNING", HARD_WARNING),
            ("WARNING", WARNING),
        ):
            found = self.by_severity(severity)
            if not found:
                continue
            lines.append(f"\n{label} ({len(found)}):")
            for issue in found[:25]:
                lines.append(f"  - {issue}")
            if len(found) > 25:
                lines.append(f"  ... and {len(found) - 25} more")
        if self.unknown_enum_values:
            lines.append("\nUnknown enum values (V2 — recorded per DD):")
            for column, values in sorted(self.unknown_enum_values.items()):
                lines.append(f"  {column}: {sorted(set(values))}")
        return "\n".join(lines)


# ─── Reading ────────────────────────────────────────────────────────────────


def read_rows(path: str | Path) -> Tuple[List[Dict[str, str]], List[str]]:
    """Read a CSV or XLSX into dicts of strings. Returns (rows, header)."""
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"corpus file not found: {path}")

    if path.suffix.lower() in (".xlsx", ".xlsm"):
        return _read_xlsx(path)
    return _read_csv(path)


def _read_csv(path: Path) -> Tuple[List[Dict[str, str]], List[str]]:
    # utf-8-sig strips the BOM Excel adds on export, which would otherwise turn
    # the first header into "﻿record_id" and fail the header check.
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        sample = fh.read(8192)
        fh.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(fh, dialect=dialect)
        header = [h.strip() for h in (reader.fieldnames or [])]
        rows = [
            {(k.strip() if k else k): (v.strip() if isinstance(v, str) else "")
             for k, v in row.items()}
            for row in reader
        ]
    return rows, header


def _read_xlsx(path: Path) -> Tuple[List[Dict[str, str]], List[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - depends on install
        raise RuntimeError(
            "Reading .xlsx needs openpyxl. Run: pip install openpyxl\n"
            "Or export the sheet to CSV and point CORPUS_CSV at it."
        ) from exc

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    it = ws.iter_rows(values_only=True)
    header_raw = next(it, ())
    header = [str(h).strip() if h is not None else "" for h in header_raw]
    rows: List[Dict[str, str]] = []
    for raw in it:
        if raw is None or all(c is None or str(c).strip() == "" for c in raw):
            continue
        row: Dict[str, str] = {}
        for idx, name in enumerate(header):
            value = raw[idx] if idx < len(raw) else None
            if isinstance(value, (datetime, date)):
                value = value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
            row[name] = "" if value is None else str(value).strip()
        rows.append(row)
    wb.close()
    return rows, header


# ─── Validation ─────────────────────────────────────────────────────────────


def validate(rows: Sequence[Dict[str, str]], header: Sequence[str], report: IngestReport
             ) -> List[Dict[str, str]]:
    """Run V1–V8. Returns the rows that survived row-level rejection."""
    _validate_header(header, report)
    if report.batch_rejected:
        return []

    _v1_record_ids(rows, report)
    if report.batch_rejected:
        return []

    accepted: List[Dict[str, str]] = []
    for row in rows:
        if _validate_row(row, report):
            accepted.append(row)
        else:
            report.rows_rejected += 1

    _v6_city_area(accepted, report)
    _v7_provider_urls(accepted, report)
    return accepted


def _validate_header(header: Sequence[str], report: IngestReport) -> None:
    missing = [c for c in REQUIRED_COLUMNS if c not in header]
    if missing:
        report.add(
            "header",
            REJECT_BATCH,
            f"missing column(s): {missing}. Data Dictionary specifies "
            f"{len(REQUIRED_COLUMNS)} columns.",
        )
    extra = [c for c in header if c and c not in REQUIRED_COLUMNS]
    if extra:
        # Not fatal: an extra column cannot corrupt a load, and rejecting the
        # batch over one would be unhelpfully brittle.
        report.add("header", WARNING, f"unexpected column(s) ignored: {extra}")


def _v1_record_ids(rows: Sequence[Dict[str, str]], report: IngestReport) -> None:
    """V1: record_id unique and matching ^MYTC-\\d{3}$ — else reject the batch."""
    seen: Dict[str, int] = {}
    bad_pattern: List[str] = []
    duplicates: List[str] = []
    for idx, row in enumerate(rows, start=2):  # +2: 1-based, and skip the header
        rid = (row.get("record_id") or "").strip()
        if not _RECORD_ID.match(rid):
            bad_pattern.append(f"line {idx}: {rid!r}")
        if rid in seen:
            duplicates.append(f"{rid} (lines {seen[rid]} and {idx})")
        else:
            seen[rid] = idx
    if bad_pattern:
        report.add(
            "V1",
            REJECT_BATCH,
            f"record_id must match ^MYTC-\\d{{3}}$ — {len(bad_pattern)} "
            f"violation(s): {bad_pattern[:5]}",
        )
    if duplicates:
        report.add(
            "V1",
            REJECT_BATCH,
            f"record_id must be unique — {len(duplicates)} duplicate(s): "
            f"{duplicates[:5]}",
        )


def _validate_row(row: Dict[str, str], report: IngestReport) -> bool:
    rid = row.get("record_id", "")
    keep = True

    # V5 first: an empty mandatory field makes the other checks meaningless.
    for column in V5_REQUIRED_NON_EMPTY:
        if not (row.get(column) or "").strip():
            report.add("V5", REJECT_ROW, "required FR-12 field is empty", rid, column)
            keep = False

    # V2: closed vocabulary, recording the foreign value as the DD asks.
    for column, allowed in ENUM_COLUMNS.items():
        value = (row.get(column) or "").strip()
        if not value:
            continue
        if value not in allowed:
            report.unknown_enum_values.setdefault(column, []).append(value)
            report.add("V2", REJECT_ROW, f"value {value!r} outside closed list", rid, column)
            keep = False

    # education_stage must agree with the level it is paired to. Not a numbered
    # DD rule, but the two columns are redundant and a disagreement would make
    # `education_stage` filters silently wrong.
    level, stage = (row.get("level") or "").strip(), (row.get("education_stage") or "").strip()
    if level in LEVEL_STAGE and stage and LEVEL_STAGE[level] != stage:
        report.add(
            "V2",
            WARNING,
            f"education_stage {stage!r} disagrees with level {level!r} "
            f"(expected {LEVEL_STAGE[level]!r})",
            rid,
            "education_stage",
        )

    # V3: https URIs.
    for column in ("provider_source_url", "curriculum_source_url"):
        value = (row.get(column) or "").strip()
        if not value:
            continue
        if not _HTTPS_URI.match(value):
            report.add("V3", REJECT_ROW, f"not a valid https URI: {value!r}", rid, column)
            keep = False

    # V4: ISO date, not in the future.
    verified = (row.get("verified_on") or "").strip()
    if verified:
        if not _ISO_DATE.match(verified):
            report.add("V4", REJECT_ROW, f"verified_on not ISO YYYY-MM-DD: {verified!r}", rid, "verified_on")
            keep = False
        else:
            try:
                parsed = date.fromisoformat(verified)
            except ValueError:
                report.add("V4", REJECT_ROW, f"verified_on unparseable: {verified!r}", rid, "verified_on")
                keep = False
            else:
                if parsed > date.today():
                    report.add(
                        "V4", REJECT_ROW,
                        f"verified_on is in the future: {verified}", rid, "verified_on",
                    )
                    keep = False

    # V8: inverted — a populated value is the alarm.
    for column in ("fee_rm", "schedule"):
        if (row.get(column) or "").strip():
            report.add(
                "V8",
                HARD_WARNING,
                f"{column} is populated. The v1 API contract declares it null "
                f"(FR-16/FR-17); a value here means the contract needs a "
                f"conscious review. Loaded as NULL to keep the contract intact.",
                rid,
                column,
            )

    return keep


def _v6_city_area(rows: Sequence[Dict[str, str]], report: IngestReport) -> None:
    """V6: city_area is bound 1:1 to state. Warning + manual review."""
    mapping: Dict[str, set] = {}
    for row in rows:
        city = (row.get("city_area") or "").strip()
        state = (row.get("state") or "").strip()
        if city and state:
            mapping.setdefault(city, set()).add(state)
    for city, states in sorted(mapping.items()):
        if len(states) > 1:
            report.add(
                "V6", WARNING,
                f"city_area {city!r} appears under multiple states: {sorted(states)}",
                column="city_area",
            )


def _v7_provider_urls(rows: Sequence[Dict[str, str]], report: IngestReport) -> None:
    """V7: provider_name maps to exactly one provider_source_url."""
    mapping: Dict[str, set] = {}
    for row in rows:
        name = (row.get("provider_name") or "").strip()
        url = (row.get("provider_source_url") or "").strip()
        if name and url:
            mapping.setdefault(name, set()).add(url)
    for name, urls in sorted(mapping.items()):
        if len(urls) > 1:
            report.add(
                "V7", WARNING,
                f"provider {name!r} maps to {len(urls)} URLs: {sorted(urls)}",
                column="provider_source_url",
            )


# ─── Seeding dimension tables ───────────────────────────────────────────────


def seed_vocabulary(session: "Session") -> Dict[str, int]:
    """Populate the enum dimension tables from app.core.vocabulary.

    Seeding from code rather than from the CSV is what makes V2 enforceable by
    the database: a foreign key can only point at a value the Data Dictionary
    sanctions, so an unknown enum cannot enter even by a direct SQL insert.
    """
    from app.infrastructure.models import (
        Curriculum,
        DeliveryMode,
        InstructionLanguage,
        Level,
        Subject,
    )

    counts: Dict[str, int] = {}

    existing = {s.name for s in session.query(Subject).all()}
    new = [Subject(name=n) for n in SUBJECTS if n not in existing]
    session.add_all(new)
    counts["subjects"] = len(new)

    existing = {l.name for l in session.query(Level).all()}
    new = [
        Level(name=n, stage=LEVEL_STAGE[n], sort_order=LEVEL_SORT_ORDER[n])
        for n in LEVELS
        if n not in existing
    ]
    session.add_all(new)
    counts["levels"] = len(new)

    existing = {m.name for m in session.query(DeliveryMode).all()}
    new = [
        DeliveryMode(name=n, mode_group=MODE_TO_GROUP[n])
        for n in DELIVERY_MODES
        if n not in existing
    ]
    session.add_all(new)
    counts["delivery_modes"] = len(new)

    existing = {i.name for i in session.query(InstructionLanguage).all()}
    new = [InstructionLanguage(name=n) for n in INSTRUCTION_LANGUAGES if n not in existing]
    session.add_all(new)
    counts["instruction_languages"] = len(new)

    existing = {c.name for c in session.query(Curriculum).all()}
    new = [Curriculum(name=n) for n in CURRICULA if n not in existing]
    session.add_all(new)
    counts["curricula"] = len(new)

    session.flush()
    return counts


def mirror_synonyms(session: "Session", entries: Iterable[Tuple[str, str, str]]) -> int:
    """Refresh the `synonyms` table from data/synonyms.json.

    The JSON stays the source of truth (NFR-H needs it editable without a
    redeploy); this mirror exists only so the dictionary is auditable in SQL.
    """
    from app.infrastructure.models import Synonym

    session.query(Synonym).delete()
    session.flush()
    rows = [
        Synonym(term=term, facet=facet, canonical_value=value)
        for term, facet, value in entries
    ]
    session.add_all(rows)
    session.flush()
    return len(rows)


# ─── Loading ────────────────────────────────────────────────────────────────


def load_rows(
    session: Session, rows: Sequence[Dict[str, str]], report: IngestReport
) -> None:
    """Insert providers then class records. Assumes `rows` already validated."""
    from app.infrastructure.models import (
        ClassRecord,
        Curriculum,
        DeliveryMode,
        InstructionLanguage,
        Level,
        Provider,
        Subject,
    )

    seed_vocabulary(session)

    subject_ids = {s.name: s.subject_id for s in session.query(Subject).all()}
    level_ids = {l.name: l.level_id for l in session.query(Level).all()}
    mode_ids = {m.name: m.delivery_mode_id for m in session.query(DeliveryMode).all()}
    lang_ids = {
        i.name: i.instruction_lang_id for i in session.query(InstructionLanguage).all()
    }
    curriculum_rows = {c.name: c for c in session.query(Curriculum).all()}

    # Providers are derived from the corpus (13 values, 1:1 with url/state/city).
    providers: Dict[str, Provider] = {p.provider_name: p for p in session.query(Provider).all()}
    for row in rows:
        name = row["provider_name"]
        if name in providers:
            continue
        provider = Provider(
            provider_name=name,
            provider_source_url=row["provider_source_url"],
            state=row["state"],
            city_area=row["city_area"],
        )
        session.add(provider)
        providers[name] = provider

    # curriculum_source_url is internal (DD: "Jangan ditampilkan sebagai sumber
    # kelas — bisa menyesatkan"), attached to the dimension row, not the class.
    for row in rows:
        curriculum = curriculum_rows.get(row.get("curriculum", ""))
        url = (row.get("curriculum_source_url") or "").strip()
        if curriculum is not None and url and not curriculum.curriculum_source_url:
            curriculum.curriculum_source_url = url

    session.flush()

    existing_ids = {r[0] for r in session.query(ClassRecord.record_id).all()}
    loaded = 0
    for row in rows:
        if row["record_id"] in existing_ids:
            report.add("load", WARNING, "record already present, skipped", row["record_id"])
            continue
        session.add(
            ClassRecord(
                record_id=row["record_id"],
                class_title=row["class_title"],
                provider_id=providers[row["provider_name"]].provider_id,
                curriculum_id=curriculum_rows[row["curriculum"]].curriculum_id,
                subject_id=subject_ids[row["subject"]],
                level_id=level_ids[row["level"]],
                delivery_mode_id=mode_ids[row["delivery_mode"]],
                instruction_lang_id=lang_ids[row["instruction_language"]],
                school_category=row["school_category"],
                education_stage=row["education_stage"],
                # Forced NULL regardless of what the file said. V8 already
                # raised a hard warning; the API contract wins until someone
                # decides otherwise.
                fee_rm=None,
                schedule=None,
                class_status=row.get("class_status") or "example_not_live_listing",
                data_label=row.get("data_label") or "web-grounded synthetic example",
                verified_on=date.fromisoformat(row["verified_on"]),
                retrieval_text=row.get("retrieval_text") or None,
                keywords=row.get("keywords") or None,
                source_basis=row.get("source_basis") or None,
            )
        )
        loaded += 1

    session.flush()
    report.rows_loaded = loaded


def ingest(session: Session, path: str | Path) -> IngestReport:
    """Full pipeline: read -> validate -> load. Nothing is written on V1 failure."""
    report = IngestReport(source=str(path))
    rows, header = read_rows(path)
    report.rows_read = len(rows)

    accepted = validate(rows, header, report)
    if report.batch_rejected:
        return report
    if not accepted:
        report.add("load", REJECT_BATCH, "no rows survived validation")
        return report

    load_rows(session, accepted, report)
    return report
